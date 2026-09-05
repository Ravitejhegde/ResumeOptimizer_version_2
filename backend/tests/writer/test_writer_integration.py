from pathlib import Path

from docx import Document

from app.analyzer.document.document_analyzer import (
    DocumentAnalyzer,
)
from app.optimizer.models.optimization_result import (
    OptimizationResult,
)
from app.optimizer.models.paragraph_update import (
    ParagraphUpdate,
)
from app.writer.engine.writer_engine import (
    WriterEngine,
)


RESUME = "tests/resources/sample_resume.docx"
OUTPUT = "tests/resources/test_writer_integration_output.docx"


def test_writer_applies_optimization_to_original_docx():

    output_path = Path(OUTPUT)

    if output_path.exists():
        output_path.unlink()

    document = DocumentAnalyzer().analyze(
        RESUME,
    )

    original_text = document.paragraphs[31]

    optimized_text = (
        "Tools & Platforms: "
        "Git, GitHub, Docker, "
        "Streamlit, Firebase, Postman"
    )

    update = ParagraphUpdate(
        paragraph_id="p31",
        section="Tools & Platforms",
        original_text=original_text,
        optimized_text=optimized_text,
        confidence=0.95,
        reason="Added user-selected skill Docker.",
        approved=True,
        formatting_safe=True,
    )

    optimization = OptimizationResult(
        document=document,
        paragraph_updates=[update],
        success=True,
        message="Optimization completed successfully.",
    )

    result = WriterEngine().write(
        source_file=RESUME,
        output_file=OUTPUT,
        document=document,
        optimization=optimization,
    )

    assert result.success is True
    assert result.output_path == OUTPUT
    assert Path(OUTPUT).exists()

    output_document = Document(OUTPUT)

    assert (
        output_document.paragraphs[31].text
        == optimized_text
    )


def test_writer_rejects_failed_optimization():

    output_path = Path(OUTPUT)

    if output_path.exists():
        output_path.unlink()

    document = DocumentAnalyzer().analyze(
        RESUME,
    )

    optimization = OptimizationResult(
        document=document,
        paragraph_updates=[],
        success=False,
        message="AI response validation failed.",
    )

    result = WriterEngine().write(
        source_file=RESUME,
        output_file=OUTPUT,
        document=document,
        optimization=optimization,
    )

    assert result.success is False
    assert result.output_path == ""
    assert not Path(OUTPUT).exists()