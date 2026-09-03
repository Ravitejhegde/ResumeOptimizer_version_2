"""
Writer integration tests.
"""

from pathlib import Path

from docx import Document

from app.analyzer.document.document_analyzer import (
    DocumentAnalyzer,
)

from app.optimizer.models.optimization_result import (
    OptimizationResult,
)

from app.optimizer.models.paragraph_update import (
    ParagraphUpdate,
)

from app.writer.services.writer import Writer


SOURCE = Path(
    r"tests\resources\sample_resume.docx"
)

OUTPUT = Path(
    r"tests\output\optimized_resume.docx"
)


def test_writer_integration() -> None:
    document = DocumentAnalyzer().analyze(SOURCE)

    optimization = OptimizationResult(
        document=document,
        success=True,
        message="Writer test",
    )

    writer = Writer()

    result = writer.write(
        source_file=SOURCE,
        output_file=OUTPUT,
        document=document,
        optimization=optimization,
    )

    assert result.success is True
    assert result.output_path is not None
    assert Path(result.output_path).exists()


def test_writer_preserves_run_formatting() -> None:
    document = DocumentAnalyzer().analyze(SOURCE)

    source_doc = Document(SOURCE)

    original_paragraph = source_doc.paragraphs[9]

    original_formatting = [
        (
            run.bold,
            run.italic,
            run.underline,
        )
        for run in original_paragraph.runs
        if run.text
    ]

    optimized_text = (
        "Edutainer PAT Technologies,"
        "AI & ML intern"
        "\t"
        "02/2026 – 6/2026 | Bengaluru"
    )

    optimization = OptimizationResult(
        document=document,
        success=True,
        message="Formatting preservation test",
        paragraph_updates=[
            ParagraphUpdate(
                paragraph_id="p9",
                section="experience",
                original_text=original_paragraph.text,
                optimized_text=optimized_text,
            )
        ],
    )

    output = Path(
        r"tests\output\format_preservation_test.docx"
    )

    writer = Writer()

    result = writer.write(
        source_file=SOURCE,
        output_file=output,
        document=document,
        optimization=optimization,
    )

    assert result.success is True
    assert output.exists()

    written_doc = Document(output)

    written_paragraph = written_doc.paragraphs[9]

    # 1. Verify that the optimized text was actually written.
    assert written_paragraph.text == optimized_text

    # 2. Verify that the paragraph still contains runs.
    written_runs = [
        run
        for run in written_paragraph.runs
        if run.text
    ]

    assert written_runs

    # 3. Verify that the original formatting pattern
    #    is preserved.
    written_formatting = [
        (
            run.bold,
            run.italic,
            run.underline,
        )
        for run in written_runs
    ]

    assert written_formatting == original_formatting

    # 4. Verify that multiple formatting runs still exist.
    assert len(written_runs) == len(original_formatting)