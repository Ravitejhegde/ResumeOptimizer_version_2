from __future__ import annotations

from copy import deepcopy

from docx import Document as DocxDocument

from app.engine.models.optimizer.optimization_result import (
    OptimizationResult,
)
from app.engine.writer.classifiers.run_classifier import (
    RunClassifier,
)
from app.engine.writer.models.paragraph_mapping import (
    ParagraphMapping,
)
from app.engine.writer.models.run_mapping import (
    RunMapping,
)
from app.engine.writer.models.write_context import (
    WriteContext,
)


class DocumentBuilder:
    """
    Builds the complete Writer context.

    No text is modified here.
    No formatting is modified here.
    No DOCX is written here.
    """

    def __init__(self) -> None:
        self._classifier = RunClassifier()

    def build(
        self,
        result: OptimizationResult,
        source_file: str,
    ) -> WriteContext:

        source_doc = DocxDocument(source_file)
        working_doc = deepcopy(source_doc)

        context = WriteContext(
            source_doc=source_doc,
            working_doc=working_doc,
            source_model=result.document,
            working_model=deepcopy(result.document),
            optimization=result,
        )

        for paragraph_index, (
            model_paragraph,
            docx_paragraph,
        ) in enumerate(
            zip(
                context.working_model.paragraphs,
                context.working_doc.paragraphs,
            )
        ):

            paragraph_mapping = ParagraphMapping(
                model_paragraph=model_paragraph,
                docx_paragraph=docx_paragraph,
                paragraph_index=paragraph_index,
            )

            for run_index, (
                model_run,
                docx_run,
            ) in enumerate(
                zip(
                    model_paragraph.runs,
                    docx_paragraph.runs,
                )
            ):

                mapping = RunMapping(
                    model_run=model_run,
                    docx_run=docx_run,
                    paragraph_index=paragraph_index,
                    run_index=run_index,
                )

                self._classifier.classify(
    mapping,
    context.working_doc,
)

                paragraph_mapping.run_mappings.append(mapping)

            context.paragraph_mappings.append(
                paragraph_mapping
            )

        return context