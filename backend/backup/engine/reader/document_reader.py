from __future__ import annotations

import logging

from docx import Document as DocxDocument

from app.engine.models.document.document import (
    Document,
)
from app.engine.reader.metadata_reader import (
    MetadataReader,
)
from app.engine.reader.page_geometry_reader import (
    PageGeometryReader,
)
from app.engine.reader.paragraph_reader import (
    ParagraphReader,
)
from app.engine.reader.table_reader import (
    TableReader,
)


logger = logging.getLogger(__name__)


class DocumentReader:
    """
    Converts DOCX document into ResumeOptimizer
    canonical Document model.

    Pipeline:

        DOCX
          |
          v
    MetadataReader
          |
          v
    ParagraphReader
          |
          v
    TableReader
          |
          v
      Document Model


    Responsibilities:
        - Read DOCX structure.
        - Preserve formatting.
        - Build engine objects.

    Does not:
        - Analyze resume.
        - Optimize content.
        - Write DOCX.
    """


    @staticmethod
    def read(
        file_path: str,
    ) -> Document:


        docx_document = DocxDocument(
            file_path
        )


        # ----------------------------------
        # Validate sections
        # ----------------------------------

        if not docx_document.sections:

            raise ValueError(
                "The document contains no sections."
            )


        # ----------------------------------
        # Metadata
        # ----------------------------------

        metadata = MetadataReader.read(
            document=docx_document,
            source_path=file_path,
        )


        document = Document(

            source_path=(
                metadata["source_path"]
            ),

            section_count=(
                len(docx_document.sections)
            ),
        )


        # ----------------------------------
        # Page geometry map
        # ----------------------------------

        section_geometry = [
            PageGeometryReader.read(
                section
            )
            for section in docx_document.sections
        ]


        default_geometry = (
            section_geometry[0]
        )


        # ----------------------------------
        # Paragraphs
        # ----------------------------------

        for index, paragraph in enumerate(
            docx_document.paragraphs,
            start=1,
        ):

            document.paragraphs.append(

                ParagraphReader.read(

                    docx_paragraph=paragraph,

                    paragraph_id=(
                        f"P{index:05}"
                    ),

                    page_geometry=(
                        default_geometry
                    ),

                    page_number=1,
                )
            )


        # ----------------------------------
        # Tables
        # ----------------------------------

        for index, table in enumerate(
            docx_document.tables,
            start=1,
        ):

            document.tables.append(

                TableReader.read(

                    docx_table=table,

                    table_id=(
                        f"T{index:03}"
                    ),
                )
            )


        logger.info(

            "DOCX loaded: paragraphs=%s tables=%s sections=%s",

            len(document.paragraphs),

            len(document.tables),

            document.section_count,
        )


        return document