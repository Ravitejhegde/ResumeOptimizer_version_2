"""
app.writer.analyzers.paragraph_locator
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~

Maps optimizer paragraph IDs to paragraphs in the
original DOCX paragraph order.
"""

from __future__ import annotations

from docx.text.paragraph import Paragraph

from app.writer.models.write_context import (
    WriteContext,
)


class ParagraphLocator:
    """
    Locates DOCX paragraphs using paragraph IDs.

    Writer paragraph IDs correspond directly to the
    original DOCX paragraph indexes.

    Empty paragraphs are therefore intentionally retained
    when building the lookup map.
    """

    def locate(
        self,
        context: WriteContext,
    ) -> None:
        """
        Build the paragraph lookup map.

        The paragraph ID ``p9`` maps directly to
        ``working_document.paragraphs[9]``.
        """

        docx_paragraphs = (
            context.working_document.paragraphs
        )

        context.paragraph_map.clear()

        for docx_index, paragraph in enumerate(
            docx_paragraphs
        ):
            paragraph_id = f"p{docx_index}"

            context.paragraph_map[
                paragraph_id
            ] = docx_index

    def get_paragraph(
        self,
        context: WriteContext,
        paragraph_id: str,
    ) -> Paragraph | None:
        """
        Return the DOCX paragraph associated with
        the supplied paragraph ID.
        """

        paragraph_index = context.paragraph_map.get(
            paragraph_id
        )

        if paragraph_index is None:
            return None

        paragraphs = (
            context.working_document.paragraphs
        )

        if not (
            0 <= paragraph_index < len(paragraphs)
        ):
            return None

        return paragraphs[paragraph_index]