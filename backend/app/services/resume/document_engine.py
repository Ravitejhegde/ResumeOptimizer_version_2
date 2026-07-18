from pathlib import Path

from docx import Document

from app.core.exceptions import ResumeParsingError
from app.models.document import (
    DocumentNode,
    ParagraphNode,
    RunNode,
)


class DocumentEngine:

    @staticmethod
    def build(path: Path) -> DocumentNode:
        try:
            document = Document(path)

            result = DocumentNode()

            for paragraph in document.paragraphs:

                paragraph_node = ParagraphNode(
                    style=paragraph.style.name if paragraph.style else "",
                    alignment=str(paragraph.alignment),
                )

                for run in paragraph.runs:

                    paragraph_node.runs.append(
                        RunNode(
                            text=run.text,
                            bold=bool(run.bold),
                            italic=bool(run.italic),
                            underline=bool(run.underline),
                            font_name=run.font.name if run.font else None,
                            font_size=float(run.font.size.pt)
                            if run.font.size
                            else None,
                        )
                    )

                result.paragraphs.append(paragraph_node)

            return result

        except Exception as e:
            raise ResumeParsingError(
                f"Unable to build document model: {e}"
            ) from e