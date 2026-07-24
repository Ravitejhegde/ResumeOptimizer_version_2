from docx import Document

from app.services.document.testing.comparison_result import (
    ComparisonResult,
)

from app.services.document.testing.paragraph_comparator import (
    ParagraphComparator,
)

from app.services.document.testing.comparators.run_comparator import (
    RunComparator,
)

from app.services.document.testing.comparators.font_comparator import (
    FontComparator,
)

from app.services.document.testing.comparators.text_comparator import (
    TextComparator,
)

from app.services.document.testing.comparators.numbering_comparator import (
    NumberingComparator,
)

from app.services.document.testing.comparators.section_comparator import (
    SectionComparator,
)

from app.services.document.testing.comparators.table_comparator import (
    TableComparator,
)

from app.services.document.testing.comparators.header_footer_comparator import (
    HeaderFooterComparator,
)

from app.services.document.testing.comparators.image_comparator import (
    ImageComparator,
)

from app.services.document.testing.comparators.page_comparator import (
    PageComparator,
)


class DocumentComparator:

    @classmethod
    def compare(
        cls,
        original_file,
        optimized_file,
    ):

        original = Document(original_file)
        optimized = Document(optimized_file)

        result = ComparisonResult()

        # Document-level comparisons
        FontComparator.compare(original, optimized, result)
        TextComparator.compare(original, optimized, result)
        NumberingComparator.compare(original, optimized, result)
        SectionComparator.compare(original, optimized, result)
        TableComparator.compare(original, optimized, result)
        HeaderFooterComparator.compare(original, optimized, result)
        ImageComparator.compare(original, optimized, result)
        PageComparator.compare(original, optimized, result)

        # Paragraph + Run comparisons
        count = min(
            len(original.paragraphs),
            len(optimized.paragraphs),
        )

        for i in range(count):

            ParagraphComparator.compare(
                original.paragraphs[i],
                optimized.paragraphs[i],
                result,
            )

            RunComparator.compare(
                original.paragraphs[i],
                optimized.paragraphs[i],
                result,
            )

        return result