from docx import Document

from app.services.document.testing.comparison_result import (
    ComparisonResult,
)


class HeaderFooterComparator:
    """
    Compares document headers and footers.
    """

    @classmethod
    def compare(
        cls,
        original: Document,
        optimized: Document,
        result: ComparisonResult,
    ):

        count = min(
            len(original.sections),
            len(optimized.sections),
        )

        for index in range(count):

            left = original.sections[index]
            right = optimized.sections[index]

            prefix = f"Section {index + 1}"

            left_header = ""
            right_header = ""

            left_footer = ""
            right_footer = ""

            if left.header.paragraphs:
                left_header = left.header.paragraphs[0].text

            if right.header.paragraphs:
                right_header = right.header.paragraphs[0].text

            if left.footer.paragraphs:
                left_footer = left.footer.paragraphs[0].text

            if right.footer.paragraphs:
                right_footer = right.footer.paragraphs[0].text

            result.add(
                f"{prefix} Header",
                left_header,
                right_header,
            )

            result.add(
                f"{prefix} Footer",
                left_footer,
                right_footer,
            )