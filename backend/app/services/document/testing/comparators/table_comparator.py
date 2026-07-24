from docx import Document

from app.services.document.testing.comparison_result import (
    ComparisonResult,
)


class TableComparator:
    """
    Compares all document tables.
    """

    @classmethod
    def compare(
        cls,
        original: Document,
        optimized: Document,
        result: ComparisonResult,
    ):

        result.add(
            "Table Count",
            len(original.tables),
            len(optimized.tables),
        )

        count = min(
            len(original.tables),
            len(optimized.tables),
        )

        for index in range(count):

            left = original.tables[index]
            right = optimized.tables[index]

            prefix = f"Table {index+1}"

            result.add(
                f"{prefix} Rows",
                len(left.rows),
                len(right.rows),
            )

            result.add(
                f"{prefix} Columns",
                len(left.columns),
                len(right.columns),
            )