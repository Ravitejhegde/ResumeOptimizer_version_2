from docx import Document

from app.services.document.testing.comparison_result import (
    ComparisonResult,
)


class NumberingComparator:
    """
    Compares list formatting (bullets/numbering).
    """

    @classmethod
    def compare(
        cls,
        original: Document,
        optimized: Document,
        result: ComparisonResult,
    ):

        count = min(
            len(original.paragraphs),
            len(optimized.paragraphs),
        )

        for index in range(count):

            left = original.paragraphs[index]
            right = optimized.paragraphs[index]

            left_style = (
                left.style.name
                if left.style
                else ""
            )

            right_style = (
                right.style.name
                if right.style
                else ""
            )

            left_list = (
                "List" in left_style
                or "Bullet" in left_style
                or "Number" in left_style
            )

            right_list = (
                "List" in right_style
                or "Bullet" in right_style
                or "Number" in right_style
            )

            result.add(
                f"Paragraph {index+1} List",
                left_list,
                right_list,
            )