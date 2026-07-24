from docx import Document

from app.services.document.testing.comparison_result import (
    ComparisonResult,
)


class TextComparator:
    """
    Compares document text content.

    This comparator detects:

    - Missing text
    - Extra text
    - Paragraph text changes
    """

    @classmethod
    def compare(
        cls,
        original: Document,
        optimized: Document,
        result: ComparisonResult,
    ):

        # -----------------------------------------
        # Full document length
        # -----------------------------------------

        original_text = "\n".join(
            paragraph.text
            for paragraph in original.paragraphs
        )

        optimized_text = "\n".join(
            paragraph.text
            for paragraph in optimized.paragraphs
        )

        result.add(
            "Document Text Length",
            len(original_text),
            len(optimized_text),
        )

        # -----------------------------------------
        # Paragraph count
        # -----------------------------------------

        result.add(
            "Paragraph Count",
            len(original.paragraphs),
            len(optimized.paragraphs),
        )

        # -----------------------------------------
        # Paragraph text
        # -----------------------------------------

        count = min(
            len(original.paragraphs),
            len(optimized.paragraphs),
        )

        for index in range(count):

            result.add(
                f"Paragraph {index+1} Text",
                original.paragraphs[index].text,
                optimized.paragraphs[index].text,
            )