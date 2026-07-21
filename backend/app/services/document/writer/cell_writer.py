from docx.table import _Row

from app.services.document.models.cell_model import (
    CellModel,
)

from app.services.document.writer.paragraph_writer import (
    ParagraphWriter,
)


class CellWriter:
    """
    Writes CellModel objects into a table row.
    """

    @classmethod
    def write(
        cls,
        row: _Row,
        cells: list[CellModel],
    ) -> None:

        for cell_model in cells:

            if cell_model.column >= len(row.cells):
                continue

            cell = row.cells[cell_model.column]

            # Remove the default empty paragraph that
            # python-docx creates in every new cell.
            if cell.paragraphs:
                p = cell.paragraphs[0]._element
                p.getparent().remove(p)

            ParagraphWriter.write(
                cell,
                cell_model.paragraphs,
            )