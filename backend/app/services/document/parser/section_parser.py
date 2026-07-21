from docx.document import Document as DocxDocument

from app.services.document.models.section_model import (
    SectionModel,
)

from app.services.document.parser.paragraph_parser import (
    ParagraphParser,
)

from app.services.document.parser.header_parser import (
    HeaderParser,
)

from app.services.document.parser.footer_parser import (
    FooterParser,
)


class SectionParser:

    @classmethod
    def parse(
        cls,
        document: DocxDocument,
    ) -> list[SectionModel]:

        sections = []

        for index, section in enumerate(document.sections):

            model = SectionModel()

            model.index = index

            model.page_width = (
                section.page_width.pt
                if section.page_width
                else 0
            )

            model.page_height = (
                section.page_height.pt
                if section.page_height
                else 0
            )

            model.margin_top = (
                section.top_margin.pt
                if section.top_margin
                else 0
            )

            model.margin_bottom = (
                section.bottom_margin.pt
                if section.bottom_margin
                else 0
            )

            model.margin_left = (
                section.left_margin.pt
                if section.left_margin
                else 0
            )

            model.margin_right = (
                section.right_margin.pt
                if section.right_margin
                else 0
            )

            model.orientation = str(
                section.orientation
            )

            model.header = HeaderParser.parse(
                section.header
            )

            model.footer = FooterParser.parse(
                section.footer
            )

            # ---------------------------------------
            # IMPORTANT FIX
            # ---------------------------------------

            if index == 0:

                model.paragraphs = (
                    ParagraphParser.parse(
                        document.paragraphs
                    )
                )

            else:

                model.paragraphs = []

            sections.append(model)

        return sections