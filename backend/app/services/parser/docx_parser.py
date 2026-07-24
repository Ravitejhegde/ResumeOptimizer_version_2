from docx import Document

from app.services.parser.models import (
    ParagraphModel,
    ResumeDocument,
    RunModel,
    TableCellModel,
    TableModel,
)


class DocxParser:

    @staticmethod
    def parse(
        file_path: str,
    ) -> ResumeDocument:

        document = Document(file_path)

        result = ResumeDocument()

        for index, paragraph in enumerate(document.paragraphs):

            fmt = paragraph.paragraph_format
            style = paragraph.style.name

            model = ParagraphModel(
                index=index,
                text=paragraph.text,
                style=style,
                alignment=paragraph.alignment,
                left_indent=fmt.left_indent.pt if fmt.left_indent else None,
                right_indent=fmt.right_indent.pt if fmt.right_indent else None,
                first_line_indent=fmt.first_line_indent.pt if fmt.first_line_indent else None,
                space_before=fmt.space_before.pt if fmt.space_before else None,
                space_after=fmt.space_after.pt if fmt.space_after else None,
                line_spacing=fmt.line_spacing,
                is_heading=style.startswith("Heading"),
                heading_level=int(style.split()[-1]) if style.startswith("Heading") and style.split()[-1].isdigit() else None,
                page_break_before=bool(fmt.page_break_before),
                keep_together=bool(fmt.keep_together),
                keep_with_next=bool(fmt.keep_with_next),
            )

            for run in paragraph.runs:

                color = None

                if run.font.color.rgb:
                    color = str(run.font.color.rgb)

                model.runs.append(
                    RunModel(
                        text=run.text,
                        bold=bool(run.bold),
                        italic=bool(run.italic),
                        underline=bool(run.underline),
                        font_name=run.font.name,
                        font_size=run.font.size.pt if run.font.size else None,
                        color=color,
                    )
                )

            model.run_count = len(model.runs)
            model.character_count = len(model.text)

            result.total_runs += model.run_count
            result.total_characters += model.character_count

            result.paragraphs.append(model)

        for table_index, table in enumerate(document.tables):

            table_model = TableModel(index=table_index)

            for row_index, row in enumerate(table.rows):

                for col_index, cell in enumerate(row.cells):

                    table_model.cells.append(
                        TableCellModel(
                            row=row_index,
                            column=col_index,
                            text=cell.text,
                        )
                    )

            result.tables.append(table_model)

        result.total_paragraphs = len(result.paragraphs)
        result.total_tables = len(result.tables)

        result.has_tables = result.total_tables > 0
        result.has_headings = any(
            p.is_heading for p in result.paragraphs
        )

        return result