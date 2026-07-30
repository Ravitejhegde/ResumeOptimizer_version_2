from __future__ import annotations

import logging
from copy import deepcopy

from docx import Document as DocxDocument

from app.engine.models.optimizer.optimization_result import (
    OptimizationResult,
)

from app.engine.writer.classifiers.run_classifier import (
    RunClassifier,
)

from app.engine.writer.models.paragraph_mapping import (
    ParagraphMapping,
)

from app.engine.writer.models.run_mapping import (
    RunMapping,
)

from app.engine.writer.models.write_context import (
    WriteContext,
)


logger = logging.getLogger(__name__)


class DocumentBuilder:
    """
    Builds Writer execution context.

    Pipeline:

        OptimizationResult
                |
                v
        DocumentBuilder
                |
                v
        WriteContext
                |
                v
        ParagraphMapping
                |
                v
        RunMapping


    Responsibilities:

        - Load DOCX.
        - Create working copy.
        - Map paragraphs.
        - Map runs.
        - Classify runs.


    Does NOT:

        - Rewrite text.
        - Save DOCX.
        - Change formatting.
    """



    def __init__(
        self,
    ) -> None:

        self._classifier = RunClassifier()



    # --------------------------------------------------

    def build(
        self,
        result: OptimizationResult,
        source_file: str,
    ) -> WriteContext:
        """
        Build writer context.
        """


        logger.info(
            "[DocumentBuilder] Loading %s",
            source_file,
        )


        source_doc = DocxDocument(
            source_file
        )


        working_doc = deepcopy(
            source_doc
        )


        working_model = deepcopy(
            result.document
        )


        context = WriteContext(

            source_doc=source_doc,

            working_doc=working_doc,

            source_model=result.document,

            working_model=working_model,

            optimization=result,

        )


        self._build_paragraph_mappings(
            context
        )


        logger.info(

            "[DocumentBuilder] "
            "paragraph mappings=%s",

            len(
                context.paragraph_mappings
            ),

        )


        return context



    # --------------------------------------------------

    def _build_paragraph_mappings(
        self,
        context: WriteContext,
    ) -> None:
        """
        Create paragraph mappings.
        """


        model_paragraphs = (
            context.working_model.paragraphs
        )


        docx_paragraphs = (
            context.working_doc.paragraphs
        )


        if len(model_paragraphs) != len(docx_paragraphs):

            context.add_warning(

                "Paragraph count mismatch "
                f"model={len(model_paragraphs)} "
                f"docx={len(docx_paragraphs)}"

            )



        limit = min(

            len(model_paragraphs),

            len(docx_paragraphs),

        )



        for index in range(limit):


            paragraph_mapping = ParagraphMapping(

                model_paragraph=
                model_paragraphs[index],


                docx_paragraph=
                docx_paragraphs[index],


                paragraph_index=index,

            )



            self._build_run_mappings(

                paragraph_mapping,

                index,

                context.working_doc,

            )



            context.paragraph_mappings.append(
                paragraph_mapping
            )



    # --------------------------------------------------

    def _build_run_mappings(
        self,
        paragraph_mapping: ParagraphMapping,
        paragraph_index: int,
        document: DocxDocument,
    ) -> None:
        """
        Create run mappings and classify runs.
        """


        model_runs = (
            paragraph_mapping
            .model_paragraph
            .runs
        )


        docx_runs = (
            paragraph_mapping
            .docx_paragraph
            .runs
        )



        if len(model_runs) != len(docx_runs):

            paragraph_mapping.warnings.append(

                "Run count mismatch "
                f"paragraph={paragraph_index}"

            )



        limit = min(

            len(model_runs),

            len(docx_runs),

        )



        for run_index in range(limit):


            mapping = RunMapping(

                model_run=model_runs[run_index],

                docx_run=docx_runs[run_index],

                paragraph_index=paragraph_index,

                run_index=run_index,

            )



            # Classify run
            self._classifier.classify(

                mapping,

                document,

            )



            paragraph_mapping.run_mappings.append(
                mapping
            )