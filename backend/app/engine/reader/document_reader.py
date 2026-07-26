from __future__ import annotations

from docx import Document as DocxDocument

from app.engine.models.document import Document
from app.engine.reader.metadata_reader import MetadataReader
from app.engine.reader.paragraph_reader import ParagraphReader
from app.engine.reader.table_reader import TableReader


class DocumentReader:
    """
    Reads an entire DOCX document into the engine
    Document model.

    Responsible ONLY for document construction.
    """

    @staticmethod
    def read(
        file_path: str,
    ) -> Document:

        docx_document = DocxDocument(file_path)

        metadata = MetadataReader.read(
            document=docx_document,
            source_path=file_path,
        )

        document = Document(

            source_path=metadata["source_path"],

            section_count=len(
                docx_document.sections
            ),

        )

        # -----------------------------
        # Paragraphs
        # -----------------------------

        for index, paragraph in enumerate(
            docx_document.paragraphs,
            start=1,
        ):

            document.paragraphs.append(

                ParagraphReader.read(

                    docx_paragraph=paragraph,

                    paragraph_id=f"P{index:05}",

                )

            )

        # -----------------------------
        # Tables
        # -----------------------------

        for index, table in enumerate(
            docx_document.tables,
            start=1,
        ):

            document.tables.append(

                TableReader.read(

                    docx_table=table,

                    table_id=f"T{index:03}",

                )

            )

        return document