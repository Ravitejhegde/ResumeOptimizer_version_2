from docx import Document

from app.document.snapshot.document_snapshot import (
    DocumentSnapshot,
)
from app.document.snapshot.paragraph_snapshot import (
    ParagraphSnapshot,
)
from app.document.snapshot.run_snapshot import (
    RunSnapshot,
)

from app.services.hyperlink.parser.hyperlink_parser import (
    HyperlinkParser,
)


class SnapshotBuilder:

    @staticmethod
    def build(
        file_path: str,
    ) -> DocumentSnapshot:

        document = Document(file_path)

        snapshot = DocumentSnapshot()

        for paragraph in document.paragraphs:

            para = ParagraphSnapshot(

                text=paragraph.text,

                style=paragraph.style.name,

            )

            # -----------------------------------------
            # Parse Hyperlinks
            # -----------------------------------------

            para.hyperlinks = HyperlinkParser.parse(
                paragraph
            )

            # -----------------------------------------
            # Parse Runs
            # -----------------------------------------

            for run in paragraph.runs:

                color = None

                if (
                    run.font.color
                    and run.font.color.rgb
                ):
                    color = str(
                        run.font.color.rgb
                    )

                size = None

                if run.font.size:
                    size = run.font.size.pt

                para.runs.append(

                    RunSnapshot(

                        text=run.text,

                        bold=bool(run.bold),

                        italic=bool(run.italic),

                        underline=bool(
                            run.underline
                        ),

                        font_name=run.font.name,

                        font_size=size,

                        color=color,

                        hyperlink=None,

                    )

                )

            snapshot.paragraphs.append(
                para
            )

        return snapshot